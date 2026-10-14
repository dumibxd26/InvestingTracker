from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from itertools import chain
from pathlib import PurePosixPath
from typing import Literal

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

MAX_BROKER_IMPORT_BYTES = 25 * 1024 * 1024
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 60 * 1024 * 1024
MAX_ARCHIVE_MEMBERS = 40
MAX_HEADER_ROWS = 30
MAX_SOURCE_ROWS = 100_000
MAX_COLUMNS = 100
MAX_RECORDS = 100_000


class BrokerImportError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedBrokerTransaction:
    event_type: Literal["closed_position", "cash_operation"]
    source_event_id: str
    source_sheet: str
    row_number: int
    broker_reference: str | None
    event_at: datetime
    opened_at: datetime | None
    closed_at: datetime | None
    symbol: str | None
    description: str | None
    direction: str | None
    asset_type: str | None
    quantity: Decimal | None
    open_price: Decimal | None
    close_price: Decimal | None
    amount: Decimal | None
    currency_code: str
    commission: Decimal | None
    swap: Decimal | None
    taxes: Decimal | None
    reported_profit: Decimal | None
    balance_after: Decimal | None


@dataclass(frozen=True)
class ParsedBrokerReport:
    content_sha256: str
    source_format: str
    transactions: tuple[ParsedBrokerTransaction, ...]
    warnings: tuple[str, ...]


_ALIASES: dict[str, tuple[str, ...]] = {
    "broker_reference": (
        "ticket",
        "position id",
        "position",
        "order id",
        "transaction id",
        "numer pozycji",
        "id pozycji",
        "id transakcji",
    ),
    "symbol": ("symbol", "instrument", "ticker", "instrument name", "nazwa instrumentu"),
    "direction": (
        "direction",
        "side",
        "buy sell",
        "type",
        "kierunek",
        "typ",
        "cumparare",
        "vanzare",
    ),
    "asset_type": (
        "instrument type",
        "asset type",
        "product type",
        "instrument category",
        "typ instrumentu",
        "rodzaj instrumentu",
    ),
    "opened_at": (
        "open time",
        "opening time",
        "open date",
        "opening date",
        "open date time",
        "data otwarcia",
        "data deschiderii",
        "czas otwarcia",
        "data otwarcia pozycji",
    ),
    "closed_at": (
        "close time",
        "closing time",
        "close date",
        "closing date",
        "close date time",
        "data zamkniecia",
        "data inchiderii",
        "czas zamkniecia",
        "data zamkniecia pozycji",
    ),
    "event_at": (
        "date",
        "date and time",
        "date time",
        "operation date",
        "operation date time",
        "operation time",
        "time",
        "transaction date",
        "data",
        "data operacji",
        "czas operacji",
    ),
    "quantity": (
        "volume",
        "quantity",
        "units",
        "shares",
        "qty",
        "wolumen",
        "ilosc",
        "cantitate",
    ),
    "open_price": (
        "open price",
        "opening price",
        "price open",
        "cena otwarcia",
        "pret deschidere",
    ),
    "close_price": (
        "close price",
        "closing price",
        "price close",
        "cena zamkniecia",
        "pret inchidere",
    ),
    "amount": ("amount", "operation amount", "value", "kwota", "wartosc", "suma"),
    "currency_code": (
        "currency",
        "profit currency",
        "currency profit",
        "account currency",
        "waluta",
        "waluta zysku",
    ),
    "commission": ("commission", "fees", "fee", "prowizja", "commission total"),
    "swap": ("swap", "financing", "rollover"),
    "taxes": ("tax", "taxes", "withholding tax", "podatek", "podatki"),
    "reported_profit": (
        "profit",
        "net profit",
        "gross profit",
        "profit loss",
        "net p l",
        "p l",
        "wynik",
        "zysk",
    ),
    "balance_after": ("balance", "balance after operation", "balance after", "saldo"),
    "description": (
        "description",
        "type",
        "operation",
        "operation type",
        "transaction type",
        "comment",
        "opis",
        "typ",
        "rodzaj operacji",
    ),
}


def _normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", text.lower())).strip()


_NORMALIZED_ALIASES = {
    field: frozenset(_normalize_header(alias) for alias in aliases)
    for field, aliases in _ALIASES.items()
}


def _header_map(row: Sequence[object]) -> dict[str, int]:
    normalized = [_normalize_header(value) for value in row]
    mapping: dict[str, int] = {}
    for field, aliases in _NORMALIZED_ALIASES.items():
        for index, value in enumerate(normalized):
            if value in aliases or any(value.startswith(f"{alias} ") for alias in aliases):
                mapping[field] = index
                break
    return mapping


def _report_kind(mapping: dict[str, int]) -> Literal["closed_position", "cash_operation"] | None:
    if {
        "symbol",
        "opened_at",
        "closed_at",
        "quantity",
        "open_price",
        "close_price",
    }.issubset(mapping):
        return "closed_position"
    if {"event_at", "amount"}.issubset(mapping):
        return "cash_operation"
    return None


def _cell(row: Sequence[object], mapping: dict[str, int], field: str) -> object | None:
    index = mapping.get(field)
    return row[index] if index is not None and index < len(row) else None


def _decimal(value: object | None, field: str, row_number: int) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    if isinstance(value, bool):
        raise BrokerImportError(f"Row {row_number}: {field} is not a valid number.")
    if isinstance(value, Decimal):
        result = value
    elif isinstance(value, (int, float)):
        result = Decimal(str(value))
    else:
        raw = str(value).strip().replace("\u00a0", "").replace(" ", "")
        negative = raw.startswith("(") and raw.endswith(")")
        if negative:
            raw = raw[1:-1]
        raw = raw.replace("'", "")
        if "," in raw and "." in raw:
            if raw.rfind(",") > raw.rfind("."):
                raw = raw.replace(".", "").replace(",", ".")
            else:
                raw = raw.replace(",", "")
        elif "," in raw:
            decimals = len(raw) - raw.rfind(",") - 1
            raw = raw.replace(",", ".") if decimals in (1, 2) else raw.replace(",", "")
        try:
            result = Decimal(raw)
        except InvalidOperation as error:
            raise BrokerImportError(
                f"Row {row_number}: {field} value {value!r} is not a valid number."
            ) from error
        if negative:
            result = -result
    if not result.is_finite():
        raise BrokerImportError(f"Row {row_number}: {field} must be finite.")
    return result


_DATE_FORMATS = (
    "%Y.%m.%d %H:%M:%S",
    "%Y.%m.%d %H:%M",
    "%d.%m.%Y %H:%M:%S",
    "%d.%m.%Y %H:%M",
    "%d.%m.%Y",
    "%d/%m/%Y %H:%M:%S",
    "%d/%m/%Y %H:%M",
    "%d/%m/%Y",
    "%m/%d/%Y %H:%M:%S",
    "%m/%d/%Y %H:%M",
    "%m/%d/%Y",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d %H:%M",
    "%Y-%m-%d",
)


def _datetime(value: object | None, field: str, row_number: int) -> datetime:
    if isinstance(value, datetime):
        return value.replace(tzinfo=None)
    if isinstance(value, date):
        return datetime.combine(value, time.min)
    if value is not None:
        text = str(value).strip()
        if text:
            try:
                parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
                return parsed.replace(tzinfo=None)
            except ValueError:
                for format_ in _DATE_FORMATS:
                    try:
                        return datetime.strptime(text, format_)
                    except ValueError:
                        continue
    raise BrokerImportError(f"Row {row_number}: {field} is missing or has an unsupported date.")


def _normalized_direction(value: object | None) -> str | None:
    normalized = _normalize_header(value)
    if normalized in {"buy", "long", "purchase", "kupno", "kup", "cumparare"}:
        return "BUY"
    if normalized in {"sell", "short", "sale", "sprzedaz", "vanzare"}:
        return "SELL"
    return str(value).strip().upper() if value is not None and str(value).strip() else None


def _currency(value: object | None, account_currency: str, row_number: int) -> str:
    code = str(value or account_currency).strip().upper()
    if not re.fullmatch(r"[A-Z]{3}", code):
        raise BrokerImportError(f"Row {row_number}: unsupported currency code {code!r}.")
    return code


def _row_event_id(
    event_type: str,
    reference: object | None,
    fields: Sequence[object | None],
    duplicate_ordinal: int,
) -> str:
    stable_reference = str(reference).strip() if reference is not None else ""
    normalized = json.dumps(
        [event_type, *[str(value) if value is not None else None for value in fields]],
        ensure_ascii=True,
        separators=(",", ":"),
    )
    digest = hashlib.sha256(normalized.encode("utf-8")).hexdigest()
    if stable_reference:
        return f"{stable_reference[:30]}:{digest[:32]}:{duplicate_ordinal}"
    return f"row:{digest}:{duplicate_ordinal}"


def _is_data_row(row: Sequence[object]) -> bool:
    return any(value is not None and str(value).strip() for value in row)


def _parse_table(
    *,
    name: str,
    rows: Iterable[Sequence[object]],
    account_currency: str,
    duplicate_ordinals: dict[str, int],
) -> tuple[list[ParsedBrokerTransaction], int, bool]:
    iterator = iter(rows)
    prefix: list[Sequence[object]] = []
    for _ in range(MAX_HEADER_ROWS):
        try:
            prefix.append(next(iterator))
        except StopIteration:
            break
    header_position = None
    mapping: dict[str, int] = {}
    event_type: Literal["closed_position", "cash_operation"] | None = None
    for index, row in enumerate(prefix):
        candidate = _header_map(row[:MAX_COLUMNS])
        kind = _report_kind(candidate)
        if kind is not None:
            header_position = index
            mapping = candidate
            event_type = kind
            break
    if header_position is None or event_type is None:
        return [], 0, False

    transactions: list[ParsedBrokerTransaction] = []
    skipped = 0
    remaining_rows = chain(prefix[header_position + 1 :], iterator)
    for row_number, row in enumerate(remaining_rows, start=header_position + 2):
        if row_number > MAX_SOURCE_ROWS:
            raise BrokerImportError(
                f"{name}: report exceeds the {MAX_SOURCE_ROWS:,}-row safety limit."
            )
        if len(row) > MAX_COLUMNS:
            raise BrokerImportError(f"{name}, row {row_number}: too many columns in report.")
        if not _is_data_row(row):
            continue
        if event_type == "closed_position":
            symbol = str(_cell(row, mapping, "symbol") or "").strip()
            if not symbol:
                skipped += 1
                continue
            opened_at = _datetime(_cell(row, mapping, "opened_at"), "open date", row_number)
            closed_at = _datetime(_cell(row, mapping, "closed_at"), "close date", row_number)
            if closed_at < opened_at:
                raise BrokerImportError(
                    f"{name}, row {row_number}: close date is before open date."
                )
            quantity = _decimal(_cell(row, mapping, "quantity"), "quantity", row_number)
            open_price = _decimal(_cell(row, mapping, "open_price"), "open price", row_number)
            close_price = _decimal(_cell(row, mapping, "close_price"), "close price", row_number)
            if quantity is None or quantity <= 0:
                raise BrokerImportError(f"{name}, row {row_number}: quantity must be positive.")
            if open_price is None or close_price is None or open_price <= 0 or close_price <= 0:
                raise BrokerImportError(f"{name}, row {row_number}: prices must be positive.")
            direction = _normalized_direction(_cell(row, mapping, "direction"))
            reported_profit = _decimal(_cell(row, mapping, "reported_profit"), "profit", row_number)
            commission = _decimal(_cell(row, mapping, "commission"), "commission", row_number)
            swap = _decimal(_cell(row, mapping, "swap"), "swap", row_number)
            taxes = _decimal(_cell(row, mapping, "taxes"), "taxes", row_number)
            currency_code = _currency(
                _cell(row, mapping, "currency_code"), account_currency, row_number
            )
            asset_type = str(_cell(row, mapping, "asset_type") or "").strip() or None
            description = None
            event_at = closed_at
            fields = (
                symbol,
                direction,
                opened_at.isoformat(),
                closed_at.isoformat(),
                str(quantity),
                str(open_price),
                str(close_price),
            )
            reference = _cell(row, mapping, "broker_reference")
            signature = hashlib.sha256(
                json.dumps(fields, separators=(",", ":")).encode("utf-8")
            ).hexdigest()
            duplicate_ordinal = duplicate_ordinals.get(signature, 0)
            duplicate_ordinals[signature] = duplicate_ordinal + 1
            source_event_id = _row_event_id("closed_position", reference, fields, duplicate_ordinal)
            amount = balance_after = None
        else:
            event_at = _datetime(_cell(row, mapping, "event_at"), "operation date", row_number)
            amount = _decimal(_cell(row, mapping, "amount"), "amount", row_number)
            if amount is None:
                raise BrokerImportError(f"{name}, row {row_number}: amount is missing.")
            description = str(_cell(row, mapping, "description") or "").strip() or None
            currency_code = _currency(
                _cell(row, mapping, "currency_code"), account_currency, row_number
            )
            balance_after = _decimal(
                _cell(row, mapping, "balance_after"), "balance after", row_number
            )
            taxes = _decimal(_cell(row, mapping, "taxes"), "taxes", row_number)
            commission = _decimal(_cell(row, mapping, "commission"), "commission", row_number)
            reference = _cell(row, mapping, "broker_reference")
            fields = (
                event_at.isoformat(),
                description,
                str(amount),
                currency_code,
                str(balance_after) if balance_after is not None else None,
            )
            signature = hashlib.sha256(
                json.dumps(fields, separators=(",", ":")).encode("utf-8")
            ).hexdigest()
            duplicate_ordinal = duplicate_ordinals.get(signature, 0)
            duplicate_ordinals[signature] = duplicate_ordinal + 1
            source_event_id = _row_event_id("cash_operation", reference, fields, duplicate_ordinal)
            opened_at = closed_at = None
            symbol = direction = asset_type = None
            quantity = open_price = close_price = swap = reported_profit = None

        transactions.append(
            ParsedBrokerTransaction(
                event_type=event_type,
                source_event_id=source_event_id,
                source_sheet=name[:80],
                row_number=row_number,
                broker_reference=(str(reference).strip()[:100] if reference is not None else None),
                event_at=event_at,
                opened_at=opened_at,
                closed_at=closed_at,
                symbol=symbol[:40] if symbol else None,
                description=description[:240] if description else None,
                direction=direction[:16] if direction else None,
                asset_type=asset_type[:40] if asset_type else None,
                quantity=quantity,
                open_price=open_price,
                close_price=close_price,
                amount=amount,
                currency_code=currency_code,
                commission=commission,
                swap=swap,
                taxes=taxes,
                reported_profit=reported_profit,
                balance_after=balance_after,
            )
        )
        if len(transactions) > MAX_RECORDS:
            raise BrokerImportError(
                f"{name}: report exceeds the {MAX_RECORDS:,} transaction safety limit."
            )
    return transactions, skipped, True


def _decode_csv(data: bytes, name: str) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1250", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise BrokerImportError(f"{name}: CSV encoding is not supported.")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(io.StringIO(text), dialect) if row]


def _parse_xlsx(
    data: bytes,
    source_name: str,
    account_currency: str,
    duplicate_ordinals: dict[str, int],
) -> tuple[list[ParsedBrokerTransaction], list[str]]:
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError) as error:
        raise BrokerImportError(f"{source_name}: workbook could not be opened as XLSX.") from error
    transactions: list[ParsedBrokerTransaction] = []
    warnings: list[str] = []
    try:
        for worksheet in workbook.worksheets:
            parsed, skipped, recognized = _parse_table(
                name=worksheet.title,
                rows=worksheet.iter_rows(values_only=True),
                account_currency=account_currency,
                duplicate_ordinals=duplicate_ordinals,
            )
            if parsed:
                transactions.extend(parsed)
            if skipped:
                warnings.append(f"{worksheet.title}: skipped {skipped} blank/summary row(s).")
            if not recognized and worksheet.max_row > 0:
                warnings.append(
                    f"{worksheet.title}: ignored; it is not a supported closed-positions or "
                    "cash-operations sheet."
                )
    finally:
        workbook.close()
    return transactions, warnings


def parse_xtb_report(data: bytes, filename: str, account_currency: str) -> ParsedBrokerReport:
    if not data:
        raise BrokerImportError("The selected report file is empty.")
    if len(data) > MAX_BROKER_IMPORT_BYTES:
        raise BrokerImportError(
            f"The report exceeds the {MAX_BROKER_IMPORT_BYTES // (1024 * 1024)} MB upload limit."
        )
    suffix = PurePosixPath(filename.replace("\\", "/")).suffix.lower()
    if suffix not in {".xlsx", ".csv", ".zip"}:
        raise BrokerImportError("Upload an XTB .xlsx, .csv, or .zip report.")

    content_sha256 = hashlib.sha256(data).hexdigest()
    duplicate_ordinals: dict[str, int] = {}
    transactions: list[ParsedBrokerTransaction] = []
    warnings: list[str] = []
    if suffix == ".xlsx":
        transactions, warnings = _parse_xlsx(data, filename, account_currency, duplicate_ordinals)
        source_format = "XTB_XLSX"
    elif suffix == ".csv":
        transactions, skipped, recognized = _parse_table(
            name=filename,
            rows=_decode_csv(data, filename),
            account_currency=account_currency,
            duplicate_ordinals=duplicate_ordinals,
        )
        source_format = "XTB_CSV"
        if skipped:
            warnings.append(f"Skipped {skipped} blank/summary row(s).")
        if not recognized:
            raise BrokerImportError(
                "CSV headers were not recognized. Export Closed Positions or Cash Operations "
                "from XTB xStation."
            )
    else:
        if not zipfile.is_zipfile(io.BytesIO(data)):
            raise BrokerImportError("The selected .zip file is not a valid ZIP archive.")
        try:
            archive = zipfile.ZipFile(io.BytesIO(data))
        except zipfile.BadZipFile as error:
            raise BrokerImportError("The selected .zip file is invalid.") from error
        with archive:
            files = [item for item in archive.infolist() if not item.is_dir()]
            if len(files) > MAX_ARCHIVE_MEMBERS:
                raise BrokerImportError("The report archive contains too many files.")
            uncompressed_size = sum(item.file_size for item in files)
            if uncompressed_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                raise BrokerImportError("The report archive expands beyond the safety limit.")
            report_files = [
                item
                for item in files
                if PurePosixPath(item.filename).suffix.lower() in {".xlsx", ".csv"}
            ]
            if not report_files:
                raise BrokerImportError(
                    "The ZIP archive contains no .xlsx or .csv transaction report."
                )
            for member in report_files:
                member_data = archive.read(member)
                member_name = PurePosixPath(member.filename).name
                if PurePosixPath(member_name).suffix.lower() == ".xlsx":
                    found, member_warnings = _parse_xlsx(
                        member_data,
                        member_name,
                        account_currency,
                        duplicate_ordinals,
                    )
                else:
                    found, skipped, recognized = _parse_table(
                        name=member_name,
                        rows=_decode_csv(member_data, member_name),
                        account_currency=account_currency,
                        duplicate_ordinals=duplicate_ordinals,
                    )
                    member_warnings = []
                    if skipped:
                        member_warnings.append(
                            f"{member_name}: skipped {skipped} blank/summary row(s)."
                        )
                    if not recognized:
                        member_warnings.append(
                            f"{member_name}: ignored because its report columns were not "
                            "recognized."
                        )
                transactions.extend(found)
                warnings.extend(member_warnings)
        source_format = "XTB_ZIP"

    if not transactions:
        details = " ".join(warnings[:3])
        raise BrokerImportError(
            "No closed positions or cash operations were found in this file. "
            "Export those report sections from XTB xStation." + (f" {details}" if details else "")
        )
    if len(transactions) > MAX_RECORDS:
        raise BrokerImportError(f"The report exceeds the {MAX_RECORDS:,} transaction safety limit.")
    return ParsedBrokerReport(
        content_sha256=content_sha256,
        source_format=source_format,
        transactions=tuple(transactions),
        warnings=tuple(warnings),
    )
