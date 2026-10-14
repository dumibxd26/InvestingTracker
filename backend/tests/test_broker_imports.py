from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.broker_imports import BrokerImportError, parse_xtb_report


def _xtb_workbook() -> bytes:
    workbook = Workbook()
    closed = workbook.active
    closed.title = "Closed Positions"
    closed.append(
        [
            "Ticket",
            "Instrument",
            "Type",
            "Open time",
            "Close time",
            "Volume",
            "Open price",
            "Close price",
            "Instrument type",
            "Currency",
            "Commission",
            "Swap",
            "Taxes",
            "Profit",
        ]
    )
    closed.append(
        [
            41234,
            "AAPL.US",
            "Buy",
            datetime(2025, 6, 1, 10, 30),
            datetime(2026, 7, 1, 14, 15),
            2,
            100,
            120,
            "Stock",
            "EUR",
            -1,
            0,
            0,
            39,
        ]
    )
    cash = workbook.create_sheet("Cash Operations")
    cash.append(["Date", "Type", "Amount", "Currency", "Balance", "Reference"])
    cash.append(
        [
            datetime(2026, 2, 2, 9, 0),
            "Dividend",
            "2,50",
            "EUR",
            "1002,50",
            "DIV-22",
        ]
    )
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def test_parse_xtb_xlsx_closed_positions_and_cash_operations() -> None:
    report = parse_xtb_report(_xtb_workbook(), "history.xlsx", "EUR")

    assert report.source_format == "XTB_XLSX"
    assert len(report.transactions) == 2
    position, cash_operation = report.transactions
    assert position.event_type == "closed_position"
    assert position.source_event_id.startswith("41234:")
    assert position.symbol == "AAPL.US"
    assert position.direction == "BUY"
    assert position.asset_type == "Stock"
    assert position.quantity == 2
    assert position.reported_profit == 39
    assert position.commission == -1
    assert cash_operation.event_type == "cash_operation"
    assert cash_operation.description == "Dividend"
    assert cash_operation.amount == 2.5
    assert cash_operation.balance_after == 1002.5
    assert cash_operation.currency_code == "EUR"


def test_parse_xtb_zip_archive_containing_excel_report() -> None:
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("reports/history.xlsx", _xtb_workbook())

    report = parse_xtb_report(stream.getvalue(), "xtb-export.zip", "EUR")

    assert report.source_format == "XTB_ZIP"
    assert len(report.transactions) == 2


def test_parser_rejects_unknown_report_columns_without_importing_partial_data() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Closed Positions"
    worksheet.append(["Ticket", "Symbol", "Profit"])
    worksheet.append([1, "AAPL.US", 100])
    stream = BytesIO()
    workbook.save(stream)

    with pytest.raises(BrokerImportError, match="No closed positions or cash operations"):
        parse_xtb_report(stream.getvalue(), "unsupported.xlsx", "EUR")
